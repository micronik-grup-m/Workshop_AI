from pptx import Presentation

from generate_deck import generate


def test_generate_en_produces_thirteen_slide_pptx():
    output_path = generate("en")

    assert output_path.exists()
    prs = Presentation(str(output_path))
    assert len(prs.slides) == 13
    for slide in prs.slides:
        has_visible_text = any(
            shape.has_text_frame and shape.text_frame.text.strip()
            for shape in slide.shapes
        )
        assert has_visible_text


def test_generate_fr_produces_thirteen_slide_pptx():
    output_path = generate("fr")

    assert output_path.exists()
    prs = Presentation(str(output_path))
    assert len(prs.slides) == 13
    for slide in prs.slides:
        has_visible_text = any(
            shape.has_text_frame and shape.text_frame.text.strip()
            for shape in slide.shapes
        )
        assert has_visible_text


def test_en_and_fr_decks_have_the_same_slide_count():
    en_prs = Presentation(str(generate("en")))
    fr_prs = Presentation(str(generate("fr")))
    assert len(en_prs.slides) == len(fr_prs.slides)
