import re
import zipfile

from pptx.enum.shapes import MSO_SHAPE_TYPE

from layouts import new_presentation
from architecture import add_architecture_diagram_slide
from generate_deck import generate


def all_text(slide):
    return "\n".join(
        shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
    )


def test_architecture_slide_has_three_boxes_and_two_arrow_labels():
    prs = new_presentation()
    slide = add_architecture_diagram_slide(
        prs, "Architecture",
        frontend_label="Angular",
        backend_label="Spring Boot",
        lmstudio_label="LM Studio",
        rest_label="REST + JWT",
        mcp_label="MCP",
    )

    text = all_text(slide)
    assert "Angular" in text
    assert "Spring Boot" in text
    assert "LM Studio" in text
    assert "REST + JWT" in text
    assert "MCP" in text


def test_architecture_slide_has_two_connectors():
    prs = new_presentation()
    slide = add_architecture_diagram_slide(
        prs, "Architecture",
        frontend_label="Angular",
        backend_label="Spring Boot",
        lmstudio_label="LM Studio",
        rest_label="REST + JWT",
        mcp_label="MCP",
    )

    connectors = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE]
    assert len(connectors) == 2


def test_architecture_slide_has_no_negative_dimensions():
    prs = new_presentation()
    slide = add_architecture_diagram_slide(
        prs, "Architecture",
        frontend_label="Angular",
        backend_label="Spring Boot",
        lmstudio_label="LM Studio",
        rest_label="REST + JWT",
        mcp_label="MCP",
    )

    for shape in slide.shapes:
        # Check width for all shapes.
        assert shape.width >= 0, f"Shape {shape.name} has negative width: {shape.width}"
        # Check height for all shapes too, including connectors. Reading
        # shape.height/.width raises ValueError when the underlying XML
        # attribute is a malformed lexical form (e.g. "3566160.0" instead
        # of an integer) — that used to be silently swallowed here, which
        # is exactly what hid the float-serialization bug that made the
        # generated .pptx files fail to open in PowerPoint. Do not
        # broaden this except/pass again: a raised ValueError here is a
        # real bug, not a legitimate absence of a height value.
        height = shape.height
        assert height >= 0, f"Shape {shape.name} has negative height: {height}"


def test_multiline_box_label_styles_every_paragraph():
    """Regression test: _add_box used to style only paragraphs[0], so the
    second line of a multi-line label (e.g. "Angular 21\n(browser)") fell
    back to the default text style (wrong size, not bold, not centered)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE as _MSO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    from pptx.dml.color import RGBColor

    prs = new_presentation()
    slide = add_architecture_diagram_slide(
        prs, "Architecture",
        frontend_label="Angular 21\n(browser)",
        backend_label="Spring Boot",
        lmstudio_label="LM Studio\n(local)",
        rest_label="REST + JWT",
        mcp_label="MCP",
    )

    multiline_boxes = [
        shape for shape in slide.shapes
        if shape.has_text_frame and len(shape.text_frame.paragraphs) > 1
        and shape.shape_type != _MSO_SHAPE_TYPE.LINE
    ]
    assert len(multiline_boxes) == 2

    for shape in multiline_boxes:
        for paragraph in shape.text_frame.paragraphs:
            assert paragraph.alignment == PP_ALIGN.CENTER
            run = paragraph.runs[0]
            assert run.font.size == Pt(14)
            assert run.font.bold is True
            assert run.font.color.rgb == RGBColor(0xFF, 0xFF, 0xFF)


def _slide_xml(pptx_path, slide_number):
    with zipfile.ZipFile(pptx_path) as archive:
        return archive.read(f"ppt/slides/slide{slide_number}.xml").decode("utf-8")


def test_saved_decks_have_well_formed_geometry():
    """Regression test for the float-serialization bug in _add_arrow's
    midpoint calculation (BOX_HEIGHT / 2 instead of // 2), which produced
    XML like x="3566160.0" — a lexical form OOXML's ST_Coordinate type
    rejects, causing PowerPoint to report a repair-needed error on open.

    This checks the SAVED .pptx files' raw XML directly (not just a
    freshly-built in-memory Presentation), since the bug only manifests
    once python-pptx serializes the Length to a string in the XML tree.
    """
    for lang in ("en", "fr"):
        output_path = generate(lang)
        xml = _slide_xml(output_path, 5)  # slide 5 is the architecture diagram in both decks

        for attr in ("x", "y", "cx", "cy"):
            for match in re.finditer(rf'{attr}="(-?[\d.]+)"', xml):
                value = match.group(1)
                assert re.fullmatch(r"\d+", value), (
                    f"{lang} slide 5 has a malformed {attr} value: {value!r} "
                    f"(must be a non-negative integer with no decimal point)"
                )
