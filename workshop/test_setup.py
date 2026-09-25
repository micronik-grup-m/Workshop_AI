# workshop/test_setup.py
from pptx import Presentation


def test_can_create_and_save_a_presentation(tmp_path):
    prs = Presentation()
    output = tmp_path / "smoke.pptx"
    prs.save(str(output))

    assert output.exists()
    reopened = Presentation(str(output))
    assert len(reopened.slides) == 0
