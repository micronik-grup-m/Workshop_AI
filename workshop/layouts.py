"""Reusable slide-layout builders for the workshop deck.

Every slide type is built on the built-in "Blank" layout (index 6) and
manually positioned text boxes/shapes, rather than the default template's
title/content placeholders — this avoids depending on placeholder indices
that can vary between python-pptx template variations, and gives full
control over a consistent, simple visual style across every slide.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

TITLE_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
BODY_COLOR = RGBColor(0x33, 0x33, 0x33)
ACCENT_COLOR = RGBColor(0x2E, 0x86, 0xAB)

BLANK_LAYOUT_INDEX = 6


def new_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[BLANK_LAYOUT_INDEX])


def add_heading(slide, text, top=Inches(0.4), font_size=32, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(0.6), top, SLIDE_WIDTH - Inches(1.2), Inches(1.0))
    text_frame = box.text_frame
    text_frame.text = text
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.size = Pt(font_size)
    run.font.bold = True
    run.font.color.rgb = TITLE_COLOR
    return box


def _add_bullets(slide, left, top, width, height, bullets, font_size=20):
    box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = box.text_frame
    text_frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = f"•  {bullet}"
        run = paragraph.runs[0]
        run.font.size = Pt(font_size)
        run.font.color.rgb = BODY_COLOR
        paragraph.space_after = Pt(12)
    return box


def add_title_slide(prs, title, subtitle, presenter_line=None):
    slide = blank_slide(prs)
    add_heading(slide, title, top=Inches(2.6), font_size=40, align=PP_ALIGN.CENTER)

    box = slide.shapes.add_textbox(Inches(0.6), Inches(3.8), SLIDE_WIDTH - Inches(1.2), Inches(1.0))
    text_frame = box.text_frame
    text_frame.text = subtitle
    paragraph = text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.runs[0]
    run.font.size = Pt(20)
    run.font.color.rgb = BODY_COLOR

    if presenter_line:
        presenter_box = slide.shapes.add_textbox(Inches(0.6), Inches(4.6), SLIDE_WIDTH - Inches(1.2), Inches(0.6))
        presenter_tf = presenter_box.text_frame
        presenter_tf.text = presenter_line
        presenter_paragraph = presenter_tf.paragraphs[0]
        presenter_paragraph.alignment = PP_ALIGN.CENTER
        presenter_run = presenter_paragraph.runs[0]
        presenter_run.font.size = Pt(16)
        presenter_run.font.italic = True
        presenter_run.font.color.rgb = BODY_COLOR

    return slide


def add_bullet_slide(prs, title, bullets):
    slide = blank_slide(prs)
    add_heading(slide, title)
    _add_bullets(
        slide, Inches(0.8), Inches(1.6),
        SLIDE_WIDTH - Inches(1.6), SLIDE_HEIGHT - Inches(2.2),
        bullets,
    )
    return slide


def add_two_column_slide(prs, title, left_heading, left_bullets, right_heading, right_bullets):
    slide = blank_slide(prs)
    add_heading(slide, title)

    column_width = (SLIDE_WIDTH - Inches(2.0)) // 2
    left_x = Inches(0.8)
    right_x = left_x + column_width + Inches(0.4)

    for x, heading, bullets in (
        (left_x, left_heading, left_bullets),
        (right_x, right_heading, right_bullets),
    ):
        heading_box = slide.shapes.add_textbox(x, Inches(1.6), column_width, Inches(0.6))
        heading_box.text_frame.text = heading
        heading_run = heading_box.text_frame.paragraphs[0].runs[0]
        heading_run.font.bold = True
        heading_run.font.size = Pt(18)
        heading_run.font.color.rgb = ACCENT_COLOR

        _add_bullets(slide, x, Inches(2.3), column_width, Inches(4.5), bullets, font_size=16)

    return slide


def add_closing_slide(prs, title, subtitle):
    return add_title_slide(prs, title, subtitle)
